import io
from pathlib import Path
import openpyxl
from typing import List, Dict, Any, Optional
from datetime import datetime
from fastapi import APIRouter, Depends, Query, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse
import math
from pydantic import BaseModel, field_validator
from sqlalchemy.orm import Session
from sqlalchemy import or_

from backend.app.db.session import get_db, BASE_DIR
from backend.app.models.entities import ProductCatalog, ProcurementBatch, SalesOrder
from backend.app.api.deps import log_audit
from backend.app.services.excel_sync import export_barcode_master_workbook, export_meesho_template_workbook

router = APIRouter(prefix="/api/catalog", tags=["Product Catalog & Barcodes"])

MASTER_COLOR_CODES: Dict[str, str] = {
    'termy green': '001', 'multi colour': '002', 'tremy pink': '003',
    'yellow': '004', 'rest red': '005', 'rest pink': '006',
    'floris pink': '007', 'floris green': '008', 'rest peach': '009',
    'rani': '010', 'rest navy': '011', 'red': '012',
    'navy blue': '013', 'pink': '014', 'boots lemon': '015',
    'khaki': '016', 'mustard': '017', 'melange': '018',
    'sky blue': '019', 'baby pink': '020', 'turk': '022',
    'gray': '023', 'charcoal': '024', 'brick red': '025',
    'sea green': '026', 'peach': '027', 'lemon': '028',
    'white': '029', 'light brown': '030', 'fawn': '031',
    'maroon': '032', 'black': '033', 'green': '034',
    'purple': '035', 'blue': '036', 'orange': '037',
    'gold': '038', 'grey': '039', 'brown': '040',
    'coral': '041', 'lavendar': '042', 'lavender': '042',
    'aqua blue': '043', 'nude': '044'
}

MASTER_SIZE_CODES: Dict[str, str] = {
    '0 - 3 months': '01', '0-3m': '01',
    '3 - 6 months': '02', '3-6m': '02',
    '6 - 9 months': '03', '6-9m': '03',
    '9 - 12 months': '05', '9-12m': '05',
    '12 - 18 months': '06', '12-18m': '06',
    '18 - 24 months': '07', '18-24m': '07',
    '24 - 36 months': '08', '24-36m': '08',
    'free': '09', 'free size': '09',
    'm': '10', 'l': '11', 'xl': '12', 'xxl': '13',
    '2xl': '13', '3xl': '14', '4xl': '15', '5xl': '16',
    's': '09', 'xs': '08'
}

class ProductCreateSchema(BaseModel):
    @field_validator("*", mode="before")
    @classmethod
    def empty_str_to_none(cls, v):
        if v == "":
            return None
        return v

    style_no: str
    season: Optional[str] = "Everyday"
    category: Optional[str] = "Nighty"
    sub_category: Optional[str] = "Sleeveless"
    product_type: Optional[str] = "Square Neck"
    sub_product: Optional[str] = "SINGLE DRESS"
    fabric_composition: Optional[str] = "Woven"
    fabric_type: Optional[str] = "Woven"
    no_of_components: Optional[int] = 1
    colour: str
    sizing: Optional[str] = "XXL"
    num_size_per_set: Optional[int] = 1
    individual_barcode: Optional[str] = None
    purchase_rate: float
    profit_margin: Optional[float] = 0.18
    meesho_price: float
    wrong_return_price: Optional[float] = None
    mrp_pcs: Optional[float] = 499.0
    mrp_set: Optional[float] = 499.0
    pack_barcode: Optional[str] = None
    image_url: Optional[str] = None
    image_url_2: Optional[str] = None
    image_url_3: Optional[str] = None
    image_url_4: Optional[str] = None
    hsn_id: Optional[str] = "620821"
    gst_pct: Optional[float] = 5.0
    net_weight_gms: Optional[int] = 285
    description: Optional[str] = None
    # Meesho specific attributes
    product_name: Optional[str] = None
    inventory: Optional[int] = 10
    country_of_origin: Optional[str] = "India"
    manufacturer_name: Optional[str] = "Pegasus Creation"
    manufacturer_address: Optional[str] = "Prasanta Apartment, Check Post"
    manufacturer_pincode: Optional[str] = "700125"
    packer_name: Optional[str] = "Divine Enterprise"
    packer_address: Optional[str] = "Prasanta Apartment, Check Post"
    packer_pincode: Optional[str] = "700125"
    importer_name: Optional[str] = None
    importer_address: Optional[str] = None
    importer_pincode: Optional[str] = None
    add_ons: Optional[str] = "No Add Ons"
    fabric: Optional[str] = "Cotton"
    fit_type: Optional[str] = "Dress"
    generic_name: Optional[str] = "Maxi"
    net_quantity: Optional[str] = "1"
    bust_size: Optional[str] = "42"
    length_size: Optional[str] = "54"
    sku_id: Optional[str] = None
    brand_name: Optional[str] = None
    group_id: Optional[str] = None
    ean_upc: Optional[str] = None
    brand: Optional[str] = None
    length: Optional[str] = "Maxi"
    neck: Optional[str] = "Square Neck"
    occasion: Optional[str] = "Everyday"
    pattern: Optional[str] = "Printed"
    pockets: Optional[str] = "No Pocket"
    print_type: Optional[str] = "Botanical"
    sleeve_length: Optional[str] = "Sleeveless"
    surface_styling: Optional[str] = "Pleated Or Gathered"
    hip_size: Optional[str] = "44"
    waist_size: Optional[str] = "36"

class ProductUpdateSchema(BaseModel):
    @field_validator("*", mode="before")
    @classmethod
    def empty_str_to_none(cls, v):
        if v == "":
            return None
        return v

    season: Optional[str] = None
    category: Optional[str] = None
    sub_category: Optional[str] = None
    product_type: Optional[str] = None
    sub_product: Optional[str] = None
    fabric_composition: Optional[str] = None
    fabric_type: Optional[str] = None
    no_of_components: Optional[int] = None
    colour: Optional[str] = None
    sizing: Optional[str] = None
    num_size_per_set: Optional[int] = None
    meesho_price: Optional[float] = None
    wrong_return_price: Optional[float] = None
    mrp_pcs: Optional[float] = None
    mrp_set: Optional[float] = None
    purchase_rate: Optional[float] = None
    profit_margin: Optional[float] = None
    image_url: Optional[str] = None
    image_url_2: Optional[str] = None
    image_url_3: Optional[str] = None
    image_url_4: Optional[str] = None
    hsn_id: Optional[str] = None
    gst_pct: Optional[float] = None
    net_weight_gms: Optional[int] = None
    description: Optional[str] = None
    is_active: Optional[bool] = None
    product_name: Optional[str] = None
    inventory: Optional[int] = None
    country_of_origin: Optional[str] = None
    manufacturer_name: Optional[str] = None
    manufacturer_address: Optional[str] = None
    manufacturer_pincode: Optional[str] = None
    packer_name: Optional[str] = None
    packer_address: Optional[str] = None
    packer_pincode: Optional[str] = None
    importer_name: Optional[str] = None
    importer_address: Optional[str] = None
    importer_pincode: Optional[str] = None
    add_ons: Optional[str] = None
    fabric: Optional[str] = None
    fit_type: Optional[str] = None
    generic_name: Optional[str] = None
    net_quantity: Optional[str] = None
    bust_size: Optional[str] = None
    length_size: Optional[str] = None
    sku_id: Optional[str] = None
    brand_name: Optional[str] = None
    group_id: Optional[str] = None
    ean_upc: Optional[str] = None
    brand: Optional[str] = None
    length: Optional[str] = None
    neck: Optional[str] = None
    occasion: Optional[str] = None
    pattern: Optional[str] = None
    pockets: Optional[str] = None
    print_type: Optional[str] = None
    sleeve_length: Optional[str] = None
    surface_styling: Optional[str] = None
    hip_size: Optional[str] = None
    waist_size: Optional[str] = None

class BatchImageUpdateSchema(BaseModel):
    updates: List[Dict[str, Any]] # [{"style_no": "...", "image_url": "..."}]

@router.get("")
def list_catalog_products(
    search: Optional[str] = Query(None, description="Search style, barcode, color"),
    colour: Optional[str] = Query(None, description="Filter by color"),
    product_type: Optional[str] = Query(None, description="Filter by neck / product type"),
    missing_images: Optional[bool] = Query(None, description="Filter items missing front image"),
    db: Session = Depends(get_db),
):
    query = db.query(ProductCatalog)

    if search:
        s = f"%{search.strip()}%"
        query = query.filter(
            or_(
                ProductCatalog.style_no.ilike(s),
                ProductCatalog.individual_barcode.ilike(s),
                ProductCatalog.pack_barcode.ilike(s),
                ProductCatalog.colour.ilike(s),
                ProductCatalog.product_type.ilike(s),
            )
        )

    if colour:
        query = query.filter(ProductCatalog.colour.ilike(f"%{colour.strip()}%"))

    if product_type:
        query = query.filter(ProductCatalog.product_type.ilike(f"%{product_type.strip()}%"))

    if missing_images is True:
        query = query.filter(
            or_(
                ProductCatalog.image_url.is_(None),
                ProductCatalog.image_url == "",
            )
        )
    elif missing_images is False:
        query = query.filter(
            ProductCatalog.image_url.is_not(None),
            ProductCatalog.image_url != "",
        )

    items = query.order_by(ProductCatalog.sl_no.asc(), ProductCatalog.id.asc()).all()

    # Pre-calculate inventory per style from procurement and sales
    proc_batches = db.query(ProcurementBatch).all()
    sales = db.query(SalesOrder).all()

    inv_map = {}
    for p in proc_batches:
        st = p.style_no.strip()
        inv_map[st] = inv_map.get(st, 0) + p.inventory
    for s in sales:
        st = s.style_no.strip()
        inv_map[st] = inv_map.get(st, 0) - s.quantity_sold

    result = []
    for item in items:
        d = item.to_dict()
        d["stock_on_hand"] = max(0, inv_map.get(item.style_no.strip(), 0))
        # Flag if price has anomaly (e.g. DE26005 shortfall or wrong return diff != 22)
        expected_wrong = item.meesho_price - 22 if item.meesho_price > 22 else 0
        d["has_price_anomaly"] = (
            abs(item.wrong_return_price - expected_wrong) > 0.5
            or (item.style_no.startswith("DE26005") and item.meesho_price < 346)
        )
        result.append(d)

    return result

@router.get("/stats")
def get_catalog_stats(db: Session = Depends(get_db)):
    products = db.query(ProductCatalog).all()
    total = len(products)
    if total == 0:
        return {
            "total_skus": 0,
            "unique_colors": 0,
            "missing_images_count": 0,
            "has_image_count": 0,
            "image_readiness_pct": 0,
            "pricing_anomalies_count": 0,
            "avg_purchase_rate": 0,
            "avg_meesho_price": 0,
            "total_valuation": 0,
        }

    colors = {p.colour for p in products if p.colour}
    missing_images = sum(1 for p in products if not p.image_url or not p.image_url.strip())
    has_images = total - missing_images
    readiness_pct = round((has_images / total) * 100, 1)

    anomalies = 0
    total_val = 0.0
    for p in products:
        expected_wrong = p.meesho_price - 22 if p.meesho_price > 22 else 0
        if abs(p.wrong_return_price - expected_wrong) > 0.5 or (p.style_no.startswith("DE26005") and p.meesho_price < 346):
            anomalies += 1
        total_val += (p.purchase_rate or 0) * 10 # default 10 units each

    avg_purchase = sum(p.purchase_rate or 0 for p in products) / total
    avg_selling = sum(p.meesho_price or 0 for p in products) / total

    return {
        "total_skus": total,
        "unique_colors": len(colors),
        "missing_images_count": missing_images,
        "has_image_count": has_images,
        "image_readiness_pct": readiness_pct,
        "pricing_anomalies_count": anomalies,
        "avg_purchase_rate": round(avg_purchase, 2),
        "avg_meesho_price": round(avg_selling, 2),
        "total_valuation": round(total_val, 2),
        "colors_list": sorted(list(colors)),
    }

@router.get("/barcode/{code}")
def resolve_barcode(code: str, db: Session = Depends(get_db)):
    """Instant barcode lookup for handheld USB or camera scanners"""
    clean_code = code.strip()
    # Check individual barcode, pack barcode, or style number
    product = db.query(ProductCatalog).filter(
        or_(
            ProductCatalog.individual_barcode == clean_code,
            ProductCatalog.pack_barcode == clean_code,
            ProductCatalog.style_no.ilike(clean_code),
        )
    ).first()

    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Barcode or Style '{clean_code}' not found in Product Catalog.",
        )

    # Calculate stock on hand
    tot_purchased = sum(
        p.inventory for p in db.query(ProcurementBatch).filter(ProcurementBatch.style_no == product.style_no).all()
    )
    tot_sold = sum(
        s.quantity_sold for s in db.query(SalesOrder).filter(SalesOrder.style_no == product.style_no).all()
    )
    stock_on_hand = max(0, tot_purchased - tot_sold)

    d = product.to_dict()
    d["stock_on_hand"] = stock_on_hand
    d["is_pack_barcode"] = (clean_code == product.pack_barcode)
    return d

@router.get("/{product_id}")
def get_catalog_product(product_id: int, db: Session = Depends(get_db)):
    product = db.query(ProductCatalog).filter(ProductCatalog.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    return product.to_dict()

@router.put("/{product_id}")
def update_catalog_product(
    product_id: int,
    payload: ProductUpdateSchema,
    db: Session = Depends(get_db),
):
    product = db.query(ProductCatalog).filter(ProductCatalog.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    update_data = payload.model_dump(exclude_unset=True) if hasattr(payload, "model_dump") else payload.dict(exclude_unset=True)
    for field, val in update_data.items():
        setattr(product, field, val)

    # Auto-adjust wrong_return_price if meesho_price updated without explicit wrong_return_price
    if "meesho_price" in update_data and "wrong_return_price" not in update_data:
        if product.meesho_price and product.meesho_price > 22:
            product.wrong_return_price = product.meesho_price - 22

    db.commit()
    db.refresh(product)

    # Keep physical Excel files synchronized immediately upon edit
    try:
        export_barcode_master_workbook(db)
        export_meesho_template_workbook(db)
    except Exception as e:
        print(f"Warning syncing excel workbooks: {e}")

    log_audit(
        db,
        category="STOCK",
        action="UPDATE",
        summary=f"Updated catalog attributes for {product.style_no}",
        details=str(update_data),
    )

    return product.to_dict()

@router.post("", status_code=201)
def create_catalog_product(payload: ProductCreateSchema, db: Session = Depends(get_db)):
    """Create a new SKU / product entry and sync immediately to both Excel workbooks"""
    # Check if style already exists
    existing = db.query(ProductCatalog).filter(ProductCatalog.style_no == payload.style_no.strip()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Product with Style No '{payload.style_no}' already exists.")

    # Calculate Barcodes if not provided
    sku = payload.individual_barcode
    if not sku:
        col_code = MASTER_COLOR_CODES.get((payload.colour or "").strip().lower(), "033")
        sz_code = MASTER_SIZE_CODES.get((payload.sizing or "").strip().lower(), "13")
        sku = f"{payload.style_no.strip()}1{col_code}{sz_code}"

    pack_bc = payload.pack_barcode or f"P{sku}"

    # Normalize margin (percentage e.g. 18 -> 0.18)
    margin = payload.profit_margin if payload.profit_margin is not None else 0.18
    if margin > 1:
        margin = margin / 100.0

    pr = payload.purchase_rate
    m_price = payload.meesho_price
    if not m_price or m_price <= 0:
        # Exact Excel formula: =ROUND($O2+10+($O2*0.2)+(($O2+10+($O2*0.2))*$P2)+(($O2+10+($O2*0.2)+(($O2+10+($O2*0.2))*$P2))*0.05),0)
        t1 = pr + 10.0 + (pr * 0.20)
        t2 = t1 + (t1 * margin)
        t3 = t2 * 1.05
        m_price = float(math.floor(t3 + 0.5))

    w_price = payload.wrong_return_price
    if w_price is None or w_price <= 0:
        w_price = max(0.0, m_price - 22.0) if m_price > 22.0 else 0.0

    max_sl = db.query(ProductCatalog.sl_no).order_by(ProductCatalog.sl_no.desc()).first()
    next_sl = (max_sl[0] + 1) if max_sl and max_sl[0] else 1

    new_prod = ProductCatalog(
        sl_no=next_sl,
        season=payload.season or "Everyday",
        style_no=payload.style_no.strip(),
        category=payload.category or "Nighty",
        sub_category=payload.sub_category or "Sleeveless",
        product_type=payload.product_type or "Square Neck",
        sub_product=payload.sub_product or "SINGLE DRESS",
        fabric_composition=payload.fabric_composition or "Woven",
        fabric_type=payload.fabric_type or "Woven",
        no_of_components=payload.no_of_components or 1,
        colour=payload.colour.strip(),
        sizing=payload.sizing or "XXL",
        num_size_per_set=payload.num_size_per_set or 1,
        individual_barcode=sku.strip(),
        purchase_rate=pr,
        profit_margin=margin,
        meesho_price=m_price,
        wrong_return_price=w_price,
        mrp_pcs=payload.mrp_pcs or 499.0,
        mrp_set=payload.mrp_set or 499.0,
        pack_barcode=pack_bc.strip(),
        image_url=payload.image_url,
        image_url_2=payload.image_url_2,
        image_url_3=payload.image_url_3,
        image_url_4=payload.image_url_4,
        hsn_id=payload.hsn_id or "620821",
        gst_pct=payload.gst_pct or 5.0,
        net_weight_gms=payload.net_weight_gms or 285,
        description=payload.description or f"Divine Enterprise {payload.colour} {payload.product_type} {payload.category}",
        is_active=True,
        # Complete Meesho template attributes
        product_name=payload.product_name,
        inventory=payload.inventory or 10,
        country_of_origin=payload.country_of_origin or "India",
        manufacturer_name=payload.manufacturer_name or "Pegasus Creation",
        manufacturer_address=payload.manufacturer_address or "Prasanta Apartment, Check Post",
        manufacturer_pincode=payload.manufacturer_pincode or "700125",
        packer_name=payload.packer_name or "Divine Enterprise",
        packer_address=payload.packer_address or "Prasanta Apartment, Check Post",
        packer_pincode=payload.packer_pincode or "700125",
        importer_name=payload.importer_name,
        importer_address=payload.importer_address,
        importer_pincode=payload.importer_pincode,
        add_ons=payload.add_ons or "No Add Ons",
        fabric=payload.fabric or "Cotton",
        fit_type=payload.fit_type or "Dress",
        generic_name=payload.generic_name or "Maxi",
        net_quantity=payload.net_quantity or "1",
        bust_size=payload.bust_size or "42",
        length_size=payload.length_size or "54",
        sku_id=payload.sku_id or sku.strip(),
        brand_name=payload.brand_name,
        group_id=payload.group_id,
        ean_upc=payload.ean_upc,
        brand=payload.brand,
        length=payload.length or "Maxi",
        neck=payload.neck or payload.product_type or "Square Neck",
        occasion=payload.occasion or "Everyday",
        pattern=payload.pattern or "Printed",
        pockets=payload.pockets or "No Pocket",
        print_type=payload.print_type or "Botanical",
        sleeve_length=payload.sleeve_length or payload.sub_category or "Sleeveless",
        surface_styling=payload.surface_styling or "Pleated Or Gathered",
        hip_size=payload.hip_size or "44",
        waist_size=payload.waist_size or "36",
    )
    db.add(new_prod)
    db.commit()
    db.refresh(new_prod)

    # Sync to Excel files
    try:
        export_barcode_master_workbook(db)
        export_meesho_template_workbook(db)
    except Exception as e:
        print(f"Warning syncing excel files: {e}")

    log_audit(
        db,
        category="STOCK",
        action="CREATE",
        summary=f"Added new product {new_prod.style_no} ({new_prod.colour}) to catalog",
        details=str(payload.model_dump()),
    )

    return new_prod.to_dict()

@router.post("/upload-image")
async def upload_product_image(file: UploadFile = File(...)):
    """Upload a new product image directly to the Pic directory with automatic web thumbnail generation"""
    filename = file.filename
    ext = Path(filename).suffix.lower()
    if ext not in [".jpg", ".jpeg", ".png", ".webp"]:
        raise HTTPException(status_code=400, detail="Only JPG, PNG, and WebP images are supported")

    upload_dir = BASE_DIR / "Pic" / "uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    thumbs_upload_dir = BASE_DIR / "Pic_thumbs" / "uploads"
    thumbs_upload_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    clean_stem = Path(filename).stem.replace(" ", "_")
    saved_filename = f"{clean_stem}_{timestamp}{ext}"
    target_path = upload_dir / saved_filename

    content = await file.read()
    with open(target_path, "wb") as f:
        f.write(content)

    # Generate thumbnail
    thumb_filename = f"{clean_stem}_{timestamp}.jpg"
    thumb_path = thumbs_upload_dir / thumb_filename
    try:
        from PIL import Image
        with Image.open(target_path) as im:
            im_thumb = im.copy()
            im_thumb.thumbnail((500, 600), Image.Resampling.LANCZOS)
            if im_thumb.mode in ("RGBA", "P"):
                im_thumb = im_thumb.convert("RGB")
            im_thumb.save(thumb_path, "JPEG", quality=85, optimize=True)
    except Exception as e:
        print(f"Thumbnail generation warning: {e}")

    return {
        "status": "success",
        "image_url": f"/Pic/uploads/{saved_filename}",
        "thumbnail_url": f"/Pic_thumbs/uploads/{thumb_filename}",
        "filename": saved_filename,
    }

@router.post("/import/excel")
async def import_excel_catalog(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upload and sync an updated Barcode Master.xlsx or MeeshoTemplate.xlsx"""
    content = await file.read()
    temp_wb = io.BytesIO(content)
    try:
        wb = openpyxl.load_workbook(temp_wb, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid Excel workbook: {str(e)}")

    imported_count = 0
    if "BARCODE" in wb.sheetnames:
        ws = wb["BARCODE"]
        for r in range(2, ws.max_row + 1):
            st = ws.cell(r, 3).value
            if not st:
                continue
            st_str = str(st).strip()
            prod = db.query(ProductCatalog).filter(ProductCatalog.style_no == st_str).first()
            if prod:
                if ws.cell(r, 2).value: prod.season = str(ws.cell(r, 2).value).strip()
                if ws.cell(r, 4).value: prod.category = str(ws.cell(r, 4).value).strip()
                if ws.cell(r, 5).value: prod.sub_category = str(ws.cell(r, 5).value).strip()
                if ws.cell(r, 6).value: prod.product_type = str(ws.cell(r, 6).value).strip()
                if ws.cell(r, 7).value: prod.sub_product = str(ws.cell(r, 7).value).strip()
                if ws.cell(r, 8).value: prod.fabric_composition = str(ws.cell(r, 8).value).strip()
                if ws.cell(r, 9).value: prod.fabric_type = str(ws.cell(r, 9).value).strip()
                if ws.cell(r, 10).value: prod.no_of_components = int(ws.cell(r, 10).value)
                if ws.cell(r, 11).value: prod.colour = str(ws.cell(r, 11).value).strip()
                if ws.cell(r, 12).value: prod.sizing = str(ws.cell(r, 12).value).strip()
                if ws.cell(r, 13).value: prod.num_size_per_set = int(ws.cell(r, 13).value)
                if ws.cell(r, 15).value is not None:
                    prod.purchase_rate = float(ws.cell(r, 15).value)
                if ws.cell(r, 16).value is not None:
                    prod.profit_margin = float(ws.cell(r, 16).value)
                if ws.cell(r, 17).value is not None:
                    prod.meesho_price = float(ws.cell(r, 17).value)
                    if prod.meesho_price > 22:
                        prod.wrong_return_price = prod.meesho_price - 22
                if ws.cell(r, 18).value is not None:
                    prod.mrp_pcs = float(ws.cell(r, 18).value)
                if ws.cell(r, 19).value is not None:
                    prod.mrp_set = float(ws.cell(r, 19).value)
                if ws.cell(r, 20).value: prod.pack_barcode = str(ws.cell(r, 20).value).strip()
                imported_count += 1
        db.commit()
    elif "Nightdress-Fill this" in wb.sheetnames:
        ws = wb["Nightdress-Fill this"]
        for r in range(3, ws.max_row + 1):
            sku = ws.cell(r, 34).value
            if not sku:
                continue
            sku_str = str(sku).strip()
            prod = db.query(ProductCatalog).filter(ProductCatalog.individual_barcode == sku_str).first()
            if prod:
                if ws.cell(r, 2).value: prod.product_name = str(ws.cell(r, 2).value).strip()
                if ws.cell(r, 3).value: prod.sizing = str(ws.cell(r, 3).value).strip()
                if ws.cell(r, 4).value is not None:
                    prod.meesho_price = float(ws.cell(r, 4).value)
                if ws.cell(r, 5).value is not None:
                    try:
                        prod.wrong_return_price = float(ws.cell(r, 5).value)
                    except (ValueError, TypeError):
                        pass
                if ws.cell(r, 6).value is not None:
                    prod.mrp_pcs = float(ws.cell(r, 6).value)
                if ws.cell(r, 7).value is not None:
                    try: prod.gst_pct = float(ws.cell(r, 7).value)
                    except (ValueError, TypeError): pass
                if ws.cell(r, 8).value: prod.hsn_id = str(ws.cell(r, 8).value).strip()
                if ws.cell(r, 9).value is not None:
                    try: prod.net_weight_gms = int(ws.cell(r, 9).value)
                    except (ValueError, TypeError): pass
                if ws.cell(r, 10).value is not None:
                    try: prod.inventory = int(ws.cell(r, 10).value)
                    except (ValueError, TypeError): pass
                if ws.cell(r, 11).value: prod.country_of_origin = str(ws.cell(r, 11).value).strip()
                if ws.cell(r, 12).value: prod.manufacturer_name = str(ws.cell(r, 12).value).strip()
                if ws.cell(r, 13).value: prod.manufacturer_address = str(ws.cell(r, 13).value).strip()
                if ws.cell(r, 14).value: prod.manufacturer_pincode = str(ws.cell(r, 14).value).strip()
                if ws.cell(r, 15).value: prod.packer_name = str(ws.cell(r, 15).value).strip()
                if ws.cell(r, 16).value: prod.packer_address = str(ws.cell(r, 16).value).strip()
                if ws.cell(r, 17).value: prod.packer_pincode = str(ws.cell(r, 17).value).strip()
                if ws.cell(r, 18).value: prod.importer_name = str(ws.cell(r, 18).value).strip()
                if ws.cell(r, 19).value: prod.importer_address = str(ws.cell(r, 19).value).strip()
                if ws.cell(r, 20).value: prod.importer_pincode = str(ws.cell(r, 20).value).strip()
                if ws.cell(r, 21).value: prod.add_ons = str(ws.cell(r, 21).value).strip()
                if ws.cell(r, 22).value: prod.colour = str(ws.cell(r, 22).value).strip()
                if ws.cell(r, 23).value: prod.fabric = str(ws.cell(r, 23).value).strip()
                if ws.cell(r, 24).value: prod.fit_type = str(ws.cell(r, 24).value).strip()
                if ws.cell(r, 25).value: prod.generic_name = str(ws.cell(r, 25).value).strip()
                if ws.cell(r, 26).value: prod.net_quantity = str(ws.cell(r, 26).value).strip()
                if ws.cell(r, 27).value: prod.bust_size = str(ws.cell(r, 27).value).strip()
                if ws.cell(r, 28).value: prod.length_size = str(ws.cell(r, 28).value).strip()
                if ws.cell(r, 29).value: prod.image_url = str(ws.cell(r, 29).value).strip()
                if ws.cell(r, 30).value: prod.image_url_2 = str(ws.cell(r, 30).value).strip()
                if ws.cell(r, 31).value: prod.image_url_3 = str(ws.cell(r, 31).value).strip()
                if ws.cell(r, 32).value: prod.image_url_4 = str(ws.cell(r, 32).value).strip()
                if ws.cell(r, 35).value: prod.brand_name = str(ws.cell(r, 35).value).strip()
                if ws.cell(r, 36).value: prod.group_id = str(ws.cell(r, 36).value).strip()
                if ws.cell(r, 37).value: prod.description = str(ws.cell(r, 37).value).strip()
                if ws.cell(r, 38).value: prod.ean_upc = str(ws.cell(r, 38).value).strip()
                if ws.cell(r, 39).value: prod.brand = str(ws.cell(r, 39).value).strip()
                if ws.cell(r, 40).value: prod.length = str(ws.cell(r, 40).value).strip()
                if ws.cell(r, 41).value: prod.neck = str(ws.cell(r, 41).value).strip()
                if ws.cell(r, 42).value: prod.occasion = str(ws.cell(r, 42).value).strip()
                if ws.cell(r, 43).value: prod.pattern = str(ws.cell(r, 43).value).strip()
                if ws.cell(r, 44).value: prod.pockets = str(ws.cell(r, 44).value).strip()
                if ws.cell(r, 45).value: prod.print_type = str(ws.cell(r, 45).value).strip()
                if ws.cell(r, 46).value: prod.sleeve_length = str(ws.cell(r, 46).value).strip()
                if ws.cell(r, 47).value: prod.surface_styling = str(ws.cell(r, 47).value).strip()
                if ws.cell(r, 48).value: prod.hip_size = str(ws.cell(r, 48).value).strip()
                if ws.cell(r, 49).value: prod.waist_size = str(ws.cell(r, 49).value).strip()
                imported_count += 1
        db.commit()
    else:
        raise HTTPException(
            status_code=400,
            detail="Unrecognized Excel format. File must contain 'BARCODE' or 'Nightdress-Fill this' sheet."
        )

    export_barcode_master_workbook(db)
    export_meesho_template_workbook(db)

    log_audit(
        db,
        category="STOCK",
        action="IMPORT",
        summary=f"Imported and synchronized {imported_count} products from uploaded Excel",
        details=file.filename,
    )
    return {"status": "success", "imported_count": imported_count}

@router.get("/export/barcode-master")
def export_barcode_master_sheet(db: Session = Depends(get_db)):
    """Generates and downloads the dynamic Barcode Master Excel sheet"""
    template_path = BASE_DIR / "Barcode Master.xlsx"
    if not template_path.exists():
        raise HTTPException(status_code=500, detail="Barcode Master file not found on server.")

    wb = openpyxl.load_workbook(template_path, data_only=False)
    if "BARCODE" in wb.sheetnames:
        ws = wb["BARCODE"]
        products = db.query(ProductCatalog).all()
        prod_map = {p.style_no.strip(): p for p in products}
        for r in range(2, ws.max_row + 1):
            st = ws.cell(r, 3).value
            if st and str(st).strip() in prod_map:
                p = prod_map[str(st).strip()]
                ws.cell(r, 15, value=p.purchase_rate)
                ws.cell(r, 16, value=p.profit_margin)
                ws.cell(r, 17, value=p.meesho_price)
                ws.cell(r, 18, value=p.mrp_pcs)
                ws.cell(r, 19, value=p.mrp_set)

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    filename = f"Barcode_Master_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@router.post("/batch-images")
def batch_update_images(payload: BatchImageUpdateSchema, db: Session = Depends(get_db)):
    """Bulk update image URLs to quickly resolve Meesho upload blockers"""
    updated_count = 0
    for u in payload.updates:
        st = u.get("style_no")
        img = u.get("image_url")
        if st and img:
            prod = db.query(ProductCatalog).filter(ProductCatalog.style_no == st.strip()).first()
            if prod:
                prod.image_url = img.strip()
                updated_count += 1
    db.commit()

    log_audit(
        db,
        category="STOCK",
        action="UPDATE",
        summary=f"Batch updated {updated_count} product image URLs",
    )
    return {"status": "success", "updated_count": updated_count}

@router.get("/export/meesho")
@router.get("/export/meesho-template")
def export_meesho_bulk_sheet(db: Session = Depends(get_db)):
    """Generates and downloads a complete, valid Meesho Bulk Upload Excel template populated with live database values"""
    template_path = BASE_DIR / "Nightdress-10177-EXTERNAL-MeeshoTemplate2PricesGSTIN-Copy.xlsx"
    if not template_path.exists():
        raise HTTPException(status_code=500, detail="Base Meesho template file not found on server.")

    # Load workbook preserving all sheets (Validation Sheet, Instructions, etc.)
    wb = openpyxl.load_workbook(template_path, data_only=False)
    if "Nightdress-Fill this" not in wb.sheetnames:
        raise HTTPException(status_code=500, detail="Sheet 'Nightdress-Fill this' not found in template.")
    ws = wb["Nightdress-Fill this"]

    products = db.query(ProductCatalog).order_by(ProductCatalog.sl_no.asc()).all()
    prod_map = {p.individual_barcode.strip(): p for p in products}
    seen_skus = set()

    # Iterate over rows in Nightdress-Fill this and update with DB values
    for r in range(3, ws.max_row + 1):
        sku = ws.cell(r, 34).value  # Col 34 is SKU ID
        if sku and str(sku).strip() in prod_map:
            sku_key = str(sku).strip()
            seen_skus.add(sku_key)
            p = prod_map[sku_key]

            if p.product_name:
                ws.cell(r, 2, value=p.product_name)
            ws.cell(r, 3, value=p.sizing or "XXL")
            ws.cell(r, 4, value=p.meesho_price)
            ws.cell(r, 5, value=p.wrong_return_price)
            ws.cell(r, 6, value=p.mrp_pcs)
            ws.cell(r, 7, value=str(int(p.gst_pct)) if p.gst_pct else "5")
            ws.cell(r, 8, value=p.hsn_id or "620821")
            ws.cell(r, 9, value=p.net_weight_gms or 285)
            ws.cell(r, 10, value=p.inventory or 10)
            ws.cell(r, 11, value=p.country_of_origin or "India")
            ws.cell(r, 12, value=p.manufacturer_name or "Pegasus Creation")
            ws.cell(r, 13, value=p.manufacturer_address or "Prasanta Apartment, Check Post")
            ws.cell(r, 14, value=p.manufacturer_pincode or "700125")
            ws.cell(r, 15, value=p.packer_name or "Divine Enterprise")
            ws.cell(r, 16, value=p.packer_address or "Prasanta Apartment, Check Post")
            ws.cell(r, 17, value=p.packer_pincode or "700125")
            if p.importer_name:
                ws.cell(r, 18, value=p.importer_name)
            if p.importer_address:
                ws.cell(r, 19, value=p.importer_address)
            if p.importer_pincode:
                ws.cell(r, 20, value=p.importer_pincode)
            ws.cell(r, 21, value=p.add_ons or "No Add Ons")
            ws.cell(r, 22, value=p.colour)
            ws.cell(r, 23, value=p.fabric or "Cotton")
            ws.cell(r, 24, value=p.fit_type or "Dress")
            ws.cell(r, 25, value=p.generic_name or "Maxi")
            ws.cell(r, 26, value=p.net_quantity or "1")
            ws.cell(r, 27, value=p.bust_size or "42")
            ws.cell(r, 28, value=p.length_size or "54")
            if p.image_url:
                ws.cell(r, 29, value=p.image_url)
            if p.image_url_2:
                ws.cell(r, 30, value=p.image_url_2)
            if p.image_url_3:
                ws.cell(r, 31, value=p.image_url_3)
            if p.image_url_4:
                ws.cell(r, 32, value=p.image_url_4)
            ws.cell(r, 33, value=p.style_no)
            ws.cell(r, 34, value=p.sku_id or p.individual_barcode)
            if p.brand_name:
                ws.cell(r, 35, value=p.brand_name)
            if p.group_id:
                ws.cell(r, 36, value=p.group_id)
            if p.description:
                ws.cell(r, 37, value=p.description)
            if p.ean_upc:
                ws.cell(r, 38, value=p.ean_upc)
            if p.brand:
                ws.cell(r, 39, value=p.brand)
            ws.cell(r, 40, value=p.length or "Maxi")
            ws.cell(r, 41, value=p.neck or p.product_type or "Square Neck")
            ws.cell(r, 42, value=p.occasion or "Everyday")
            ws.cell(r, 43, value=p.pattern or "Printed")
            ws.cell(r, 44, value=p.pockets or "No Pocket")
            ws.cell(r, 45, value=p.print_type or "Botanical")
            ws.cell(r, 46, value=p.sleeve_length or p.sub_category or "Sleeveless")
            ws.cell(r, 47, value=p.surface_styling or "Pleated Or Gathered")
            ws.cell(r, 48, value=p.hip_size or "44")
            ws.cell(r, 49, value=p.waist_size or "36")

    # Append any new products not already in sheet
    for p in products:
        if p.individual_barcode.strip() not in seen_skus:
            ws.append([
                None,
                p.product_name,
                p.sizing or "XXL",
                p.meesho_price,
                p.wrong_return_price,
                p.mrp_pcs,
                str(int(p.gst_pct)) if p.gst_pct else "5",
                p.hsn_id or "620821",
                p.net_weight_gms or 285,
                p.inventory or 10,
                p.country_of_origin or "India",
                p.manufacturer_name or "Pegasus Creation",
                p.manufacturer_address or "Prasanta Apartment, Check Post",
                p.manufacturer_pincode or "700125",
                p.packer_name or "Divine Enterprise",
                p.packer_address or "Prasanta Apartment, Check Post",
                p.packer_pincode or "700125",
                p.importer_name,
                p.importer_address,
                p.importer_pincode,
                p.add_ons or "No Add Ons",
                p.colour,
                p.fabric or "Cotton",
                p.fit_type or "Dress",
                p.generic_name or "Maxi",
                p.net_quantity or "1",
                p.bust_size or "42",
                p.length_size or "54",
                p.image_url,
                p.image_url_2,
                p.image_url_3,
                p.image_url_4,
                p.style_no,
                p.sku_id or p.individual_barcode,
                p.brand_name,
                p.group_id,
                p.description,
                p.ean_upc,
                p.brand,
                p.length or "Maxi",
                p.neck or p.product_type or "Square Neck",
                p.occasion or "Everyday",
                p.pattern or "Printed",
                p.pockets or "No Pocket",
                p.print_type or "Botanical",
                p.sleeve_length or p.sub_category or "Sleeveless",
                p.surface_styling or "Pleated Or Gathered",
                p.hip_size or "44",
                p.waist_size or "36",
            ])

    # Save to in-memory bytes buffer
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)

    filename = f"Meesho_Catalog_Upload_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

