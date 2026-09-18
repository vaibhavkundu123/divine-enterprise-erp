import csv
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from backend.app.db.session import SessionLocal, BASE_DIR, DATA_DIR
from backend.app.models.entities import (
    ProcurementBatch,
    SalesOrder,
    RTOPipeline,
    CustomerReturn,
    ItemExchange,
    AdSpend,
    BankTransaction,
)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SUMMARY_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
SUMMARY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
BORDER_DOUBLE = Side(border_style="double", color="000000")
BORDER_TOP_THIN = Side(border_style="thin", color="000000")

TOTAL_BORDER = Border(top=BORDER_TOP_THIN, bottom=BORDER_DOUBLE, left=BORDER_THIN, right=BORDER_THIN)
CELL_BORDER = Border(top=BORDER_THIN, bottom=BORDER_THIN, left=BORDER_THIN, right=BORDER_THIN)

CURRENCY_FORMAT = "$#,##0.00"
PERCENT_FORMAT = "0.00%"
INT_FORMAT = "#,##0"

def safe_save_workbook(wb: openpyxl.Workbook, target_path: Path, max_retries: int = 3) -> Path:
    for attempt in range(max_retries):
        try:
            wb.save(target_path)
            return target_path
        except PermissionError:
            time.sleep(0.5)
    # If locked, save to fallback with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    fallback_path = target_path.parent / f"{target_path.stem}_fallback_{timestamp}.xlsx"
    wb.save(fallback_path)
    return fallback_path

def apply_header_style(ws, headers: List[str]):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER

def auto_fit_columns(ws, max_cols: Optional[int] = None):
    cols_to_fit = max_cols if max_cols else ws.max_column
    for col_idx in range(1, cols_to_fit + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def export_flat_csvs(db: Session, out_dir: Path = DATA_DIR):
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. procurement_batches.csv
    proc_items = db.query(ProcurementBatch).order_by(ProcurementBatch.date.asc()).all()
    with open(out_dir / "procurement_batches.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "date", "style_no", "inventory", "purchase_rate", "total_value"])
        for p in proc_items:
            writer.writerow([p.id, p.date, p.style_no, p.inventory, p.purchase_rate, p.total_value])

    # 2. sales_orders.csv
    sales = db.query(SalesOrder).order_by(SalesOrder.date.asc()).all()
    with open(out_dir / "sales_orders.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "sl_no", "date", "style_no", "quantity_sold", "selling_price", "total_revenue", "unit_purchase_cost", "cogs", "profit", "profit_margin", "reference"])
        for s in sales:
            writer.writerow([s.id, s.sl_no, s.date, s.style_no, s.quantity_sold, s.selling_price, s.total_revenue, s.unit_purchase_cost, s.cogs, s.profit, s.profit_margin, s.reference])

    # 3. rto_pipeline.csv
    rtos = db.query(RTOPipeline).order_by(RTOPipeline.date.asc()).all()
    with open(out_dir / "rto_pipeline.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "rto_id", "date", "style_no", "quantity", "sale_price", "reversed_revenue", "courier_fee", "status", "received_date", "restocked_date", "tracking_no", "notes"])
        for r in rtos:
            writer.writerow([r.id, r.rto_id, r.date, r.style_no, r.quantity, r.sale_price, r.reversed_revenue, r.courier_fee, r.status, r.received_date, r.restocked_date, r.tracking_no, r.notes])

    # 4. customer_returns.csv
    returns = db.query(CustomerReturn).order_by(CustomerReturn.date.asc()).all()
    with open(out_dir / "customer_returns.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "return_id", "date", "style_no", "quantity", "refund_amount", "reverse_fee", "primary_reason", "secondary_reason", "status", "qc_grade", "received_date", "restocked_date", "reverse_awb", "notes"])
        for cr in returns:
            writer.writerow([cr.id, cr.return_id, cr.date, cr.style_no, cr.quantity, cr.refund_amount, cr.reverse_fee, cr.primary_reason, cr.secondary_reason, cr.status, cr.qc_grade, cr.received_date, cr.restocked_date, cr.reverse_awb, cr.notes])

    # 5. item_exchanges.csv
    exchanges = db.query(ItemExchange).order_by(ItemExchange.date.asc()).all()
    with open(out_dir / "item_exchanges.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "exchange_id", "date", "original_style", "exchanged_style", "quantity", "standard_price", "reverse_fee", "amount_received", "cogs", "profit", "primary_reason", "secondary_reason", "return_status", "exchange_status", "reverse_awb", "received_date", "restocked_date"])
        for ex in exchanges:
            writer.writerow([ex.id, ex.exchange_id, ex.date, ex.original_style, ex.exchanged_style, ex.quantity, ex.standard_price, ex.reverse_fee, ex.amount_received, ex.cogs, ex.profit, ex.primary_reason, ex.secondary_reason, ex.return_status, ex.exchange_status, ex.reverse_awb, ex.received_date, ex.restocked_date])

    # 6. ad_spends.csv
    ads = db.query(AdSpend).order_by(AdSpend.date.asc()).all()
    with open(out_dir / "ad_spends.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "date", "platform", "amount", "notes"])
        for a in ads:
            writer.writerow([a.id, a.date, a.platform, a.amount, a.notes])

    # 7. bank_transactions.csv
    banks = db.query(BankTransaction).order_by(BankTransaction.date.asc()).all()
    with open(out_dir / "bank_transactions.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["id", "sl_no", "date", "type", "amount", "running_balance"])
        for b in banks:
            writer.writerow([b.id, b.sl_no, b.date, b.type, b.amount, b.running_balance])

def export_logistic_workbook(db: Session, target_path: Optional[Path] = None) -> Path:
    if target_path is None:
        target_path = BASE_DIR / "Logistic.xlsx"

    wb = openpyxl.Workbook()
    # Sheet 1: Transactions
    ws_tx = wb.active
    ws_tx.title = "Transactions"
    tx_headers = ["Date", "Style No.", "Inventory", "Purchase Rate", "Total"]
    apply_header_style(ws_tx, tx_headers)

    batches = db.query(ProcurementBatch).order_by(ProcurementBatch.date.asc(), ProcurementBatch.id.asc()).all()
    daily_groups: Dict[str, Dict[str, Any]] = {}

    row_num = 2
    for b in batches:
        ws_tx.append([
            b.date,
            b.style_no,
            b.inventory,
            b.purchase_rate,
            f"=C{row_num}*D{row_num}",
        ])
        ws_tx.cell(row=row_num, column=3).number_format = INT_FORMAT
        ws_tx.cell(row=row_num, column=4).number_format = CURRENCY_FORMAT
        ws_tx.cell(row=row_num, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 6):
            ws_tx.cell(row=row_num, column=c).border = CELL_BORDER

        # Group for daily summary
        if b.date not in daily_groups:
            daily_groups[b.date] = {"styles": set(), "units": 0, "total": 0.0}
        daily_groups[b.date]["styles"].add(b.style_no)
        daily_groups[b.date]["units"] += b.inventory
        daily_groups[b.date]["total"] += b.total_value

        row_num += 1

    # Total Summary row for Transactions
    if row_num > 2:
        ws_tx.append([
            "Total Summary",
            "",
            f"=SUM(C2:C{row_num-1})",
            "",
            f"=SUM(E2:E{row_num-1})",
        ])
        for c in range(1, 6):
            cell = ws_tx.cell(row=row_num, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c in (3, 5):
                cell.number_format = INT_FORMAT if c == 3 else CURRENCY_FORMAT
    auto_fit_columns(ws_tx)

    # Sheet 2: Daily Summary
    ws_sum = wb.create_sheet(title="Daily Summary")
    sum_headers = ["Date", "Style Count", "Total Units", "Daily Gross Total", "Avg Purchase Rate"]
    apply_header_style(ws_sum, sum_headers)

    s_row = 2
    for d_key in sorted(daily_groups.keys()):
        grp = daily_groups[d_key]
        style_cnt = len(grp["styles"])
        tot_units = grp["units"]
        gross_tot = grp["total"]
        ws_sum.append([
            d_key,
            style_cnt,
            tot_units,
            gross_tot,
            f"=IF(C{s_row}>0, D{s_row}/C{s_row}, 0)",
        ])
        ws_sum.cell(row=s_row, column=2).number_format = INT_FORMAT
        ws_sum.cell(row=s_row, column=3).number_format = INT_FORMAT
        ws_sum.cell(row=s_row, column=4).number_format = CURRENCY_FORMAT
        ws_sum.cell(row=s_row, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 6):
            ws_sum.cell(row=s_row, column=c).border = CELL_BORDER
        s_row += 1

    if s_row > 2:
        ws_sum.append([
            "Total Summary",
            "",
            f"=SUM(C2:C{s_row-1})",
            f"=SUM(D2:D{s_row-1})",
            f"=IF(C{s_row}>0, D{s_row}/C{s_row}, 0)",
        ])
        for c in range(1, 6):
            cell = ws_sum.cell(row=s_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c in (3, 4, 5):
                cell.number_format = INT_FORMAT if c == 3 else CURRENCY_FORMAT
    auto_fit_columns(ws_sum)

    return safe_save_workbook(wb, target_path)

def export_sales_inventory_workbook(db: Session, target_path: Optional[Path] = None) -> Path:
    if target_path is None:
        target_path = BASE_DIR / "Sales_Inventory.xlsx"

    wb = openpyxl.Workbook()

    # 1. Sales Sheet
    ws_sales = wb.active
    ws_sales.title = "Sales"
    sales_headers = ["Sl No.", "Date", "Style No.", "Quantity Sold", "Selling Price ($)", "Total Revenue ($)", "Cost of Goods Sold ($)", "Profit ($)", "Profit Margin (%)", "Reference"]
    apply_header_style(ws_sales, sales_headers)

    sales = db.query(SalesOrder).order_by(SalesOrder.date.asc(), SalesOrder.id.asc()).all()
    row_num = 2
    for s in sales:
        ws_sales.append([
            s.sl_no or (row_num - 1),
            s.date,
            s.style_no,
            s.quantity_sold,
            s.selling_price,
            f"=D{row_num}*E{row_num}",
            s.cogs,
            f"=F{row_num}-G{row_num}",
            f"=IF(F{row_num}>0, H{row_num}/F{row_num}, 0)",
            s.reference or "Direct Sale",
        ])
        ws_sales.cell(row=row_num, column=4).number_format = INT_FORMAT
        ws_sales.cell(row=row_num, column=5).number_format = CURRENCY_FORMAT
        ws_sales.cell(row=row_num, column=6).number_format = CURRENCY_FORMAT
        ws_sales.cell(row=row_num, column=7).number_format = CURRENCY_FORMAT
        ws_sales.cell(row=row_num, column=8).number_format = CURRENCY_FORMAT
        ws_sales.cell(row=row_num, column=9).number_format = PERCENT_FORMAT
        for c in range(1, 11):
            ws_sales.cell(row=row_num, column=c).border = CELL_BORDER
        row_num += 1

    if row_num > 2:
        ws_sales.append([
            "Total Summary", "", "",
            f"=SUM(D2:D{row_num-1})",
            "",
            f"=SUM(F2:F{row_num-1})",
            f"=SUM(G2:G{row_num-1})",
            f"=SUM(H2:H{row_num-1})",
            f"=IF(F{row_num}>0, H{row_num}/F{row_num}, 0)",
            "",
        ])
        for c in range(1, 11):
            cell = ws_sales.cell(row=row_num, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 4:
                cell.number_format = INT_FORMAT
            elif c in (6, 7, 8):
                cell.number_format = CURRENCY_FORMAT
            elif c == 9:
                cell.number_format = PERCENT_FORMAT
    auto_fit_columns(ws_sales)

    # 2. Stock Inventory Sheet
    ws_stock = wb.create_sheet(title="Stock Inventory")
    stock_headers = ["Style No.", "Total Purchased", "Total Sold", "In-Transit RTO", "RTO Holding", "Restocked RTO", "Stock on Hand", "Unit Cost ($)", "Stock Valuation ($)", "Total Profit ($)", "Status"]
    apply_header_style(ws_stock, stock_headers)

    # Gather distinct styles across procurement and sales
    proc_styles = db.query(ProcurementBatch.style_no).distinct().all()
    sales_styles = db.query(SalesOrder.style_no).distinct().all()
    all_styles = sorted(list({s[0] for s in proc_styles if s[0]} | {s[0] for s in sales_styles if s[0]}))

    row_s = 2
    for st in all_styles:
        tot_purchased = sum(p.inventory for p in db.query(ProcurementBatch).filter(ProcurementBatch.style_no == st).all())
        tot_sold = sum(s.quantity_sold for s in db.query(SalesOrder).filter(SalesOrder.style_no == st).all())
        in_transit = sum(r.quantity for r in db.query(RTOPipeline).filter(RTOPipeline.style_no == st, RTOPipeline.status == "In Transit").all())
        rto_holding = sum(r.quantity for r in db.query(RTOPipeline).filter(RTOPipeline.style_no == st, RTOPipeline.status == "Received").all())
        restocked = sum(r.quantity for r in db.query(RTOPipeline).filter(RTOPipeline.style_no == st, RTOPipeline.status == "Restocked").all())
        # Also restocked customer returns
        restocked += sum(cr.quantity for cr in db.query(CustomerReturn).filter(CustomerReturn.style_no == st, CustomerReturn.status == "Restocked").all())

        # WAC Unit Cost
        p_batches = db.query(ProcurementBatch).filter(ProcurementBatch.style_no == st).all()
        tot_cost = sum(p.total_value for p in p_batches)
        tot_inv = sum(p.inventory for p in p_batches)
        wac = (tot_cost / tot_inv) if tot_inv > 0 else 0.0

        # Total Profit from sales
        st_profit = sum(s.profit for s in db.query(SalesOrder).filter(SalesOrder.style_no == st).all())

        ws_stock.append([
            st,
            tot_purchased,
            tot_sold,
            in_transit,
            rto_holding,
            restocked,
            f"=B{row_s}-C{row_s}+F{row_s}",
            wac,
            f"=G{row_s}*H{row_s}",
            st_profit,
            f'=IF(G{row_s}>10, "In Stock", IF(G{row_s}>0, "Low Stock", IF(G{row_s}=0, "Out of Stock", "Over Sold")))',
        ])
        ws_stock.cell(row=row_s, column=2).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=3).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=4).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=5).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=6).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=7).number_format = INT_FORMAT
        ws_stock.cell(row=row_s, column=8).number_format = CURRENCY_FORMAT
        ws_stock.cell(row=row_s, column=9).number_format = CURRENCY_FORMAT
        ws_stock.cell(row=row_s, column=10).number_format = CURRENCY_FORMAT
        for c in range(1, 12):
            ws_stock.cell(row=row_s, column=c).border = CELL_BORDER
        row_s += 1

    if row_s > 2:
        ws_stock.append([
            "Total Summary",
            f"=SUM(B2:B{row_s-1})",
            f"=SUM(C2:C{row_s-1})",
            f"=SUM(D2:D{row_s-1})",
            f"=SUM(E2:E{row_s-1})",
            f"=SUM(F2:F{row_s-1})",
            f"=SUM(G2:G{row_s-1})",
            "",
            f"=SUM(I2:I{row_s-1})",
            f"=SUM(J2:J{row_s-1})",
            "",
        ])
        for c in range(1, 12):
            cell = ws_stock.cell(row=row_s, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c in (2, 3, 4, 5, 6, 7):
                cell.number_format = INT_FORMAT
            elif c in (9, 10):
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_stock)

    # 3. RTO Tracker
    ws_rto = wb.create_sheet(title="RTO Tracker")
    rto_headers = ["RTO ID", "Date", "Style No.", "Quantity", "Sale Price ($)", "Reversed Revenue ($)", "Courier Fee ($)", "Status", "Received Date", "Restocked Date", "Tracking No.", "Notes"]
    apply_header_style(ws_rto, rto_headers)
    rtos = db.query(RTOPipeline).order_by(RTOPipeline.date.asc(), RTOPipeline.id.asc()).all()
    r_row = 2
    for r in rtos:
        ws_rto.append([
            r.rto_id, r.date, r.style_no, r.quantity, r.sale_price,
            f"=D{r_row}*E{r_row}", r.courier_fee, r.status,
            r.received_date or "", r.restocked_date or "", r.tracking_no or "", r.notes or ""
        ])
        ws_rto.cell(row=r_row, column=4).number_format = INT_FORMAT
        ws_rto.cell(row=r_row, column=5).number_format = CURRENCY_FORMAT
        ws_rto.cell(row=r_row, column=6).number_format = CURRENCY_FORMAT
        ws_rto.cell(row=r_row, column=7).number_format = CURRENCY_FORMAT
        for c in range(1, 13):
            ws_rto.cell(row=r_row, column=c).border = CELL_BORDER
        r_row += 1
    if r_row > 2:
        ws_rto.append([
            "Total Summary", "", "", f"=SUM(D2:D{r_row-1})", "", f"=SUM(F2:F{r_row-1})", f"=SUM(G2:G{r_row-1})",
            "", "", "", "", ""
        ])
        for c in range(1, 13):
            cell = ws_rto.cell(row=r_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 4:
                cell.number_format = INT_FORMAT
            elif c in (6, 7):
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_rto)

    # 4. Customer Returns
    ws_cr = wb.create_sheet(title="Customer Returns")
    cr_headers = ["Return ID", "Date", "Style No.", "Quantity", "Refund Amount ($)", "Reverse Fee ($)", "Primary Reason", "Secondary Reason", "Status", "Received Date", "Restocked Date", "Reverse AWB", "Notes"]
    apply_header_style(ws_cr, cr_headers)
    rets = db.query(CustomerReturn).order_by(CustomerReturn.date.asc(), CustomerReturn.id.asc()).all()
    cr_row = 2
    for cr in rets:
        ws_cr.append([
            cr.return_id, cr.date, cr.style_no, cr.quantity, cr.refund_amount, cr.reverse_fee,
            cr.primary_reason, cr.secondary_reason or "", cr.status, cr.received_date or "",
            cr.restocked_date or "", cr.reverse_awb or "", cr.notes or ""
        ])
        ws_cr.cell(row=cr_row, column=4).number_format = INT_FORMAT
        ws_cr.cell(row=cr_row, column=5).number_format = CURRENCY_FORMAT
        ws_cr.cell(row=cr_row, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 14):
            ws_cr.cell(row=cr_row, column=c).border = CELL_BORDER
        cr_row += 1
    if cr_row > 2:
        ws_cr.append([
            "Total Summary", "", "", f"=SUM(D2:D{cr_row-1})", f"=SUM(E2:E{cr_row-1})", f"=SUM(F2:F{cr_row-1})",
            "", "", "", "", "", "", ""
        ])
        for c in range(1, 14):
            cell = ws_cr.cell(row=cr_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 4:
                cell.number_format = INT_FORMAT
            elif c in (5, 6):
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_cr)

    # 5. Exchanges
    ws_ex = wb.create_sheet(title="Exchanges")
    ex_headers = ["Exchange ID", "Date", "Original Style (Returned)", "Exchanged Style (New)", "Quantity", "Standard Price ($)", "Reverse Fee ($)", "Amount Received ($)", "COGS ($)", "Profit ($)", "Primary Reason", "Secondary Reason", "Return Status", "Exchange Status", "Reverse AWB", "Received Date", "Restocked Date"]
    apply_header_style(ws_ex, ex_headers)
    exchs = db.query(ItemExchange).order_by(ItemExchange.date.asc(), ItemExchange.id.asc()).all()
    ex_row = 2
    for ex in exchs:
        ws_ex.append([
            ex.exchange_id, ex.date, ex.original_style, ex.exchanged_style, ex.quantity,
            ex.standard_price, ex.reverse_fee, f"=(E{ex_row}*F{ex_row})-G{ex_row}",
            ex.cogs, f"=H{ex_row}-I{ex_row}", ex.primary_reason or "", ex.secondary_reason or "",
            ex.return_status, ex.exchange_status, ex.reverse_awb or "", ex.received_date or "", ex.restocked_date or ""
        ])
        ws_ex.cell(row=ex_row, column=5).number_format = INT_FORMAT
        for c in (6, 7, 8, 9, 10):
            ws_ex.cell(row=ex_row, column=c).number_format = CURRENCY_FORMAT
        for c in range(1, 18):
            ws_ex.cell(row=ex_row, column=c).border = CELL_BORDER
        ex_row += 1
    if ex_row > 2:
        ws_ex.append([
            "Total Summary", "", "", "", f"=SUM(E2:E{ex_row-1})", "", f"=SUM(G2:G{ex_row-1})",
            f"=SUM(H2:H{ex_row-1})", f"=SUM(I2:I{ex_row-1})", f"=SUM(J2:J{ex_row-1})",
            "", "", "", "", "", "", ""
        ])
        for c in range(1, 18):
            cell = ws_ex.cell(row=ex_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 5:
                cell.number_format = INT_FORMAT
            elif c in (7, 8, 9, 10):
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_ex)

    # 6. Ad Spend
    ws_ad = wb.create_sheet(title="Ad Spend")
    ad_headers = ["Date", "Platform", "Amount ($)", "Notes"]
    apply_header_style(ws_ad, ad_headers)
    ads = db.query(AdSpend).order_by(AdSpend.date.asc(), AdSpend.id.asc()).all()
    ad_groups: Dict[str, Dict[str, Any]] = {}
    ad_row = 2
    for a in ads:
        ws_ad.append([a.date, a.platform, a.amount, a.notes or ""])
        ws_ad.cell(row=ad_row, column=3).number_format = CURRENCY_FORMAT
        for c in range(1, 5):
            ws_ad.cell(row=ad_row, column=c).border = CELL_BORDER
        
        if a.date not in ad_groups:
            ad_groups[a.date] = {"count": 0, "platforms": set(), "amount": 0.0}
        ad_groups[a.date]["count"] += 1
        ad_groups[a.date]["platforms"].add(a.platform)
        ad_groups[a.date]["amount"] += a.amount

        ad_row += 1
    if ad_row > 2:
        ws_ad.append(["Total Summary", "", f"=SUM(C2:C{ad_row-1})", ""])
        for c in range(1, 5):
            cell = ws_ad.cell(row=ad_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 3:
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_ad)

    # 7. Bank Transactions
    ws_bank = wb.create_sheet(title="Bank Transactions")
    bank_headers = ["Sl No.", "Date", "Type", "Amount ($)", "Running Balance ($)"]
    apply_header_style(ws_bank, bank_headers)
    banks = db.query(BankTransaction).order_by(BankTransaction.date.asc(), BankTransaction.id.asc()).all()
    b_row = 2
    for b in banks:
        ws_bank.append([b.sl_no or (b_row - 1), b.date, b.type, b.amount, b.running_balance])
        ws_bank.cell(row=b_row, column=4).number_format = CURRENCY_FORMAT
        ws_bank.cell(row=b_row, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 6):
            ws_bank.cell(row=b_row, column=c).border = CELL_BORDER
        b_row += 1
    if b_row > 2:
        ws_bank.append(["Total Summary", "", "", f"=SUM(D2:D{b_row-1})", ""])
        for c in range(1, 6):
            cell = ws_bank.cell(row=b_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 4:
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_bank)

    # 8. Daily Sales Summary
    ws_d_sales = wb.create_sheet(title="Daily Sales Summary")
    d_sales_headers = ["Date", "Orders Count", "Units Sold", "Gross Revenue ($)", "Total COGS ($)", "Gross Profit ($)", "Gross Margin (%)"]
    apply_header_style(ws_d_sales, d_sales_headers)

    daily_sales: Dict[str, Dict[str, Any]] = {}
    for s in sales:
        if s.date not in daily_sales:
            daily_sales[s.date] = {"orders": 0, "units": 0, "rev": 0.0, "cogs": 0.0}
        daily_sales[s.date]["orders"] += 1
        daily_sales[s.date]["units"] += s.quantity_sold
        daily_sales[s.date]["rev"] += s.total_revenue
        daily_sales[s.date]["cogs"] += s.cogs

    ds_row = 2
    for d_key in sorted(daily_sales.keys()):
        d_val = daily_sales[d_key]
        ws_d_sales.append([
            d_key,
            d_val["orders"],
            d_val["units"],
            d_val["rev"],
            d_val["cogs"],
            f"=D{ds_row}-E{ds_row}",
            f"=IF(D{ds_row}>0, F{ds_row}/D{ds_row}, 0)",
        ])
        ws_d_sales.cell(row=ds_row, column=2).number_format = INT_FORMAT
        ws_d_sales.cell(row=ds_row, column=3).number_format = INT_FORMAT
        ws_d_sales.cell(row=ds_row, column=4).number_format = CURRENCY_FORMAT
        ws_d_sales.cell(row=ds_row, column=5).number_format = CURRENCY_FORMAT
        ws_d_sales.cell(row=ds_row, column=6).number_format = CURRENCY_FORMAT
        ws_d_sales.cell(row=ds_row, column=7).number_format = PERCENT_FORMAT
        for c in range(1, 8):
            ws_d_sales.cell(row=ds_row, column=c).border = CELL_BORDER
        ds_row += 1

    if ds_row > 2:
        ws_d_sales.append([
            "Total Summary",
            f"=SUM(B2:B{ds_row-1})",
            f"=SUM(C2:C{ds_row-1})",
            f"=SUM(D2:D{ds_row-1})",
            f"=SUM(E2:E{ds_row-1})",
            f"=SUM(F2:F{ds_row-1})",
            f"=IF(D{ds_row}>0, F{ds_row}/D{ds_row}, 0)",
        ])
        for c in range(1, 8):
            cell = ws_d_sales.cell(row=ds_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c in (2, 3):
                cell.number_format = INT_FORMAT
            elif c in (4, 5, 6):
                cell.number_format = CURRENCY_FORMAT
            elif c == 7:
                cell.number_format = PERCENT_FORMAT
    auto_fit_columns(ws_d_sales)

    # 9. Daily Ad Spend Summary
    ws_d_ad = wb.create_sheet(title="Daily Ad Spend Summary")
    d_ad_headers = ["Date", "Campaign Count", "Platforms Used", "Total Ad Spend ($)"]
    apply_header_style(ws_d_ad, d_ad_headers)
    dad_row = 2
    for d_key in sorted(ad_groups.keys()):
        ag = ad_groups[d_key]
        plat_str = ", ".join(sorted(ag["platforms"]))
        ws_d_ad.append([d_key, ag["count"], plat_str, ag["amount"]])
        ws_d_ad.cell(row=dad_row, column=2).number_format = INT_FORMAT
        ws_d_ad.cell(row=dad_row, column=4).number_format = CURRENCY_FORMAT
        for c in range(1, 5):
            ws_d_ad.cell(row=dad_row, column=c).border = CELL_BORDER
        dad_row += 1

    if dad_row > 2:
        ws_d_ad.append(["Total Summary", f"=SUM(B2:B{dad_row-1})", "", f"=SUM(D2:D{dad_row-1})"])
        for c in range(1, 5):
            cell = ws_d_ad.cell(row=dad_row, column=c)
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = TOTAL_BORDER
            if c == 2:
                cell.number_format = INT_FORMAT
            elif c == 4:
                cell.number_format = CURRENCY_FORMAT
    auto_fit_columns(ws_d_ad)

    return safe_save_workbook(wb, target_path)

def sync_all(db: Optional[Session] = None) -> Dict[str, Any]:
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True
    try:
        log_path = export_logistic_workbook(db)
        sales_path = export_sales_inventory_workbook(db)
        export_flat_csvs(db)
        return {
            "status": "success",
            "logistic_path": str(log_path),
            "sales_inventory_path": str(sales_path),
            "timestamp": datetime.now().isoformat(),
        }
    finally:
        if close_db:
            db.close()
