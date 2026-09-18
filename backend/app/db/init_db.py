import os
import openpyxl
from datetime import datetime, date
from pathlib import Path
from sqlalchemy.orm import Session
from backend.app.db.session import engine, SessionLocal, BASE_DIR
from backend.app.models.entities import (
    Base,
    ProcurementBatch,
    SalesOrder,
    RTOPipeline,
    CustomerReturn,
    ItemExchange,
    AdSpend,
    BankTransaction,
    ActivityAuditLog,
)

def format_date(val) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    val_str = str(val).strip()
    if " 00:00:00" in val_str:
        val_str = val_str.replace(" 00:00:00", "")
    return val_str

def safe_float(val, default: float = 0.0) -> float:
    if val is None:
        return default
    try:
        if isinstance(val, str):
            val = val.replace("$", "").replace(",", "").strip()
        return float(val)
    except (ValueError, TypeError):
        return default

def safe_int(val, default: int = 0) -> int:
    if val is None:
        return default
    try:
        if isinstance(val, str):
            val = val.replace(",", "").strip()
        return int(float(val))
    except (ValueError, TypeError):
        return default

def is_summary_row(val) -> bool:
    if not val:
        return False
    val_lower = str(val).strip().lower()
    return any(marker in val_lower for marker in ["total summary", "grand total", "total"])

def init_db(db: Session = None) -> None:
    # 1. Create tables
    Base.metadata.create_all(bind=engine)
    
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True

    try:
        # Check if already seeded
        existing_proc = db.query(ProcurementBatch).first()
        existing_sales = db.query(SalesOrder).first()

        logistic_path = BASE_DIR / "Logistic.xlsx"
        sales_path = BASE_DIR / "Sales_Inventory.xlsx"

        # Ingest Logistic.xlsx
        if not existing_proc and logistic_path.exists():
            wb = openpyxl.load_workbook(logistic_path, data_only=True)
            if "Transactions" in wb.sheetnames:
                ws = wb["Transactions"]
                last_date = ""
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    date_val, style_no, inv, rate, total = row[:5] if len(row) >= 5 else (*row, *([None] * (5 - len(row))))
                    
                    if is_summary_row(date_val) or is_summary_row(style_no):
                        continue

                    d_str = format_date(date_val)
                    if d_str:
                        last_date = d_str
                    else:
                        d_str = last_date

                    if not style_no:
                        continue

                    inv_int = safe_int(inv, 0)
                    rate_flt = safe_float(rate, 0.0)
                    total_flt = safe_float(total, inv_int * rate_flt)

                    batch = ProcurementBatch(
                        date=d_str,
                        style_no=str(style_no).strip(),
                        inventory=inv_int,
                        purchase_rate=rate_flt,
                        total_value=total_flt,
                    )
                    db.add(batch)
                db.commit()

        # Ingest Sales_Inventory.xlsx
        if not existing_sales and sales_path.exists():
            wb_sales = openpyxl.load_workbook(sales_path, data_only=True)

            # 1. Sales Sheet
            if "Sales" in wb_sales.sheetnames:
                ws = wb_sales["Sales"]
                last_date = ""
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    sl_no, date_val, style_no, qty, price, rev, cogs, profit, margin, ref = (
                        row[:10] if len(row) >= 10 else (*row, *([None] * (10 - len(row))))
                    )
                    if is_summary_row(sl_no) or is_summary_row(date_val) or is_summary_row(style_no):
                        continue
                    
                    d_str = format_date(date_val)
                    if d_str:
                        last_date = d_str
                    else:
                        d_str = last_date

                    if not style_no:
                        continue

                    qty_int = safe_int(qty, 1)
                    price_flt = safe_float(price, 0.0)
                    rev_flt = safe_float(rev, qty_int * price_flt)
                    cogs_flt = safe_float(cogs, 0.0)
                    profit_flt = safe_float(profit, rev_flt - cogs_flt)
                    margin_flt = safe_float(margin, profit_flt / rev_flt if rev_flt > 0 else 0.0)
                    unit_cost = cogs_flt / qty_int if qty_int > 0 else 0.0

                    order = SalesOrder(
                        sl_no=safe_int(sl_no, None),
                        date=d_str,
                        style_no=str(style_no).strip(),
                        quantity_sold=qty_int,
                        selling_price=price_flt,
                        total_revenue=rev_flt,
                        unit_purchase_cost=round(unit_cost, 2),
                        cogs=cogs_flt,
                        profit=profit_flt,
                        profit_margin=margin_flt,
                        reference=str(ref).strip() if ref else None,
                    )
                    db.add(order)

            # 2. RTO Tracker
            if "RTO Tracker" in wb_sales.sheetnames:
                ws = wb_sales["RTO Tracker"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    rto_id, date_val, style_no, qty, price, rev, fee, status, rec_d, rest_d, track, notes = (
                        row[:12] if len(row) >= 12 else (*row, *([None] * (12 - len(row))))
                    )
                    if is_summary_row(rto_id) or is_summary_row(date_val) or not rto_id:
                        continue
                    
                    rto = RTOPipeline(
                        rto_id=str(rto_id).strip(),
                        date=format_date(date_val),
                        style_no=str(style_no).strip() if style_no else "",
                        quantity=safe_int(qty, 1),
                        sale_price=safe_float(price, 0.0),
                        reversed_revenue=safe_float(rev, 0.0),
                        courier_fee=safe_float(fee, 0.0),
                        status=str(status).strip() if status else "In Transit",
                        received_date=format_date(rec_d) if rec_d else None,
                        restocked_date=format_date(rest_d) if rest_d else None,
                        tracking_no=str(track).strip() if track else None,
                        notes=str(notes).strip() if notes else None,
                    )
                    db.add(rto)

            # 3. Customer Returns
            if "Customer Returns" in wb_sales.sheetnames:
                ws = wb_sales["Customer Returns"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    ret_id, date_val, style_no, qty, refund, fee, p_reas, s_reas, status, rec_d, rest_d, awb, *extra = (
                        row if len(row) >= 12 else (*row, *([None] * (12 - len(row))))
                    )
                    if is_summary_row(ret_id) or is_summary_row(date_val) or not ret_id:
                        continue
                    
                    ret = CustomerReturn(
                        return_id=str(ret_id).strip(),
                        date=format_date(date_val),
                        style_no=str(style_no).strip() if style_no else "",
                        quantity=safe_int(qty, 1),
                        refund_amount=safe_float(refund, 0.0),
                        reverse_fee=safe_float(fee, 175.0),
                        primary_reason=str(p_reas).strip() if p_reas else "General Return",
                        secondary_reason=str(s_reas).strip() if s_reas else None,
                        status=str(status).strip() if status else "In Transit",
                        qc_grade=None,
                        received_date=format_date(rec_d) if rec_d else None,
                        restocked_date=format_date(rest_d) if rest_d else None,
                        reverse_awb=str(awb).strip() if awb else None,
                        notes=str(extra[0]).strip() if extra and extra[0] else None,
                    )
                    db.add(ret)

            # 4. Exchanges
            if "Exchanges" in wb_sales.sheetnames:
                ws = wb_sales["Exchanges"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    ex_id, date_val, orig, exch, qty, std_p, fee, amt_r, cogs, profit, p_reas, s_reas, ret_stat, ex_stat, awb, *extra = (
                        row if len(row) >= 15 else (*row, *([None] * (15 - len(row))))
                    )
                    if is_summary_row(ex_id) or is_summary_row(date_val) or not ex_id:
                        continue
                    
                    rec_d = extra[0] if len(extra) > 0 else None
                    rest_d = extra[1] if len(extra) > 1 else None

                    exch_item = ItemExchange(
                        exchange_id=str(ex_id).strip(),
                        date=format_date(date_val),
                        original_style=str(orig).strip() if orig else "",
                        exchanged_style=str(exch).strip() if exch else "",
                        quantity=safe_int(qty, 1),
                        standard_price=safe_float(std_p, 0.0),
                        reverse_fee=safe_float(fee, 175.0),
                        amount_received=safe_float(amt_r, 0.0),
                        cogs=safe_float(cogs, 0.0),
                        profit=safe_float(profit, 0.0),
                        primary_reason=str(p_reas).strip() if p_reas else None,
                        secondary_reason=str(s_reas).strip() if s_reas else None,
                        return_status=str(ret_stat).strip() if ret_stat else "In Transit",
                        exchange_status=str(ex_stat).strip() if ex_stat else "Dispatched",
                        reverse_awb=str(awb).strip() if awb else None,
                        received_date=format_date(rec_d) if rec_d else None,
                        restocked_date=format_date(rest_d) if rest_d else None,
                    )
                    db.add(exch_item)

            # 5. Ad Spend
            if "Ad Spend" in wb_sales.sheetnames:
                ws = wb_sales["Ad Spend"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    date_val, platform, amount, notes = (
                        row[:4] if len(row) >= 4 else (*row, *([None] * (4 - len(row))))
                    )
                    if is_summary_row(date_val) or not date_val:
                        continue
                    
                    ad = AdSpend(
                        date=format_date(date_val),
                        platform=str(platform).strip() if platform else "Other",
                        amount=safe_float(amount, 0.0),
                        notes=str(notes).strip() if notes else None,
                    )
                    db.add(ad)

            # 6. Bank Transactions
            if "Bank Transactions" in wb_sales.sheetnames:
                ws = wb_sales["Bank Transactions"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not any(row):
                        continue
                    sl_no, date_val, b_type, amount, balance = (
                        row[:5] if len(row) >= 5 else (*row, *([None] * (5 - len(row))))
                    )
                    if is_summary_row(sl_no) or is_summary_row(date_val) or not date_val:
                        continue
                    
                    bank = BankTransaction(
                        sl_no=safe_int(sl_no, None),
                        date=format_date(date_val),
                        type=str(b_type).strip() if b_type else "Credited (+)",
                        amount=safe_float(amount, 0.0),
                        running_balance=safe_float(balance, 0.0),
                    )
                    db.add(bank)

            # Audit Log Initial Entry
            audit = ActivityAuditLog(
                timestamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                category="SYSTEM",
                action="INITIALIZE",
                summary="System cold-start initialized and seeded from master workbooks",
                status="SUCCESS",
                source="System Seed",
                details="Seeded initial data from Logistic.xlsx and Sales_Inventory.xlsx",
            )
            db.add(audit)
            db.commit()

    finally:
        if close_db:
            db.close()

if __name__ == "__main__":
    init_db()
    print("Database initialization and cold-start seeding completed successfully.")
