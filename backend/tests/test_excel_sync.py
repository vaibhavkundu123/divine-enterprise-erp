import openpyxl
import pytest
from pathlib import Path
from backend.app.db.session import SessionLocal, BASE_DIR, DATA_DIR
from backend.app.db.init_db import init_db
from backend.app.services.excel_sync import sync_all, export_logistic_workbook, export_sales_inventory_workbook, export_flat_csvs
from backend.app.models.entities import ProcurementBatch, SalesOrder

def test_database_initialization_and_seed():
    db = SessionLocal()
    init_db(db)
    proc_count = db.query(ProcurementBatch).count()
    sales_count = db.query(SalesOrder).count()
    db.close()
    assert proc_count > 0, "Procurement batches should be seeded"
    assert sales_count > 0, "Sales orders should be seeded"

def test_excel_sync_workbooks_and_sheets(tmp_path):
    db = SessionLocal()
    test_log_path = tmp_path / "Test_Logistic.xlsx"
    test_sales_path = tmp_path / "Test_Sales_Inventory.xlsx"

    log_path = export_logistic_workbook(db, target_path=test_log_path)
    sales_path = export_sales_inventory_workbook(db, target_path=test_sales_path)

    # Check Logistic sheets
    wb_log = openpyxl.load_workbook(log_path, data_only=False)
    assert "Transactions" in wb_log.sheetnames
    assert "Daily Summary" in wb_log.sheetnames
    ws_tx = wb_log["Transactions"]
    assert ws_tx.max_row > 2
    # Verify formulas in transactions
    formula_found = any("=" in str(cell.value) for cell in ws_tx[2] if cell.value)
    assert formula_found, "Transactions should contain formulas"

    # Check Sales_Inventory 9 sheets
    wb_sales = openpyxl.load_workbook(sales_path, data_only=False)
    expected_sheets = [
        "Sales", "Stock Inventory", "RTO Tracker", "Customer Returns",
        "Exchanges", "Ad Spend", "Bank Transactions", "Daily Sales Summary",
        "Daily Ad Spend Summary"
    ]
    for s_name in expected_sheets:
        assert s_name in wb_sales.sheetnames, f"Missing sheet {s_name} in Sales_Inventory.xlsx"

    ws_sales = wb_sales["Sales"]
    assert ws_sales.max_row > 2
    sales_formula = any("=" in str(cell.value) for cell in ws_sales[2] if cell.value)
    assert sales_formula, "Sales sheet should contain dynamic formulas"

    db.close()

def test_flat_csv_generation(tmp_path):
    db = SessionLocal()
    export_flat_csvs(db, out_dir=tmp_path)
    csv_files = [
        "procurement_batches.csv",
        "sales_orders.csv",
        "rto_pipeline.csv",
        "customer_returns.csv",
        "item_exchanges.csv",
        "ad_spends.csv",
        "bank_transactions.csv",
    ]
    for fn in csv_files:
        f_path = tmp_path / fn
        assert f_path.exists(), f"Missing CSV {fn}"
        assert f_path.stat().st_size > 0, f"CSV {fn} is empty"
    db.close()
